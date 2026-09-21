"""Read explicit tables without guessing graph relationships."""

import csv
from pathlib import Path

import pandas as pd


def read_table(path, *, sheet=None, page=None, table=None, header_row=1):
    """Return literal cell values; selectors and header rows are one-based."""
    path = Path(path)
    if header_row < 1 or (page is not None and page < 1) or (table is not None and table < 1):
        raise ValueError("header-row, page, and table must be positive integers")
    suffix = path.suffix.lower()
    try:
        if suffix in {".csv", ".tsv"}:
            if any(value is not None for value in (sheet, page, table)):
                raise ValueError("CSV/TSV files do not accept sheet, page, or table selectors")
            with path.open(encoding="utf-8-sig", newline="") as stream:
                rows = list(csv.reader(stream, delimiter="\t" if suffix == ".tsv" else ","))
        elif suffix == ".xlsx":
            from openpyxl import load_workbook

            if page is not None or table is not None:
                raise ValueError("Excel files accept a sheet selector, not page/table")
            workbook = load_workbook(path, read_only=True, data_only=False)
            try:
                if sheet is None and len(workbook.sheetnames) != 1:
                    raise ValueError(
                        f"Choose --sheet (prepare: --nodes-sheet/--edges-sheet): "
                        f"{', '.join(workbook.sheetnames)}"
                    )
                selected = sheet or workbook.sheetnames[0]
                if selected not in workbook.sheetnames:
                    raise ValueError(
                        f"Unknown sheet '{selected}'; available: {workbook.sheetnames}"
                    )
                rows = []
                for cells in workbook[selected].iter_rows():
                    if any(cell.data_type == "f" for cell in cells):
                        raise ValueError(
                            "Excel formulas must be replaced with values before import"
                        )
                    rows.append([cell.value for cell in cells])
            finally:
                workbook.close()
        elif suffix == ".pdf":
            import pdfplumber

            if sheet is not None:
                raise ValueError("PDF files accept page/table selectors, not sheet")
            with pdfplumber.open(path) as document:
                if page is not None and page > len(document.pages):
                    raise ValueError(f"PDF has only {len(document.pages)} pages")
                candidates = []
                for number, pdf_page in enumerate(document.pages, 1):
                    if page is not None and number != page:
                        continue
                    for index, rows_found in enumerate(pdf_page.extract_tables(), 1):
                        if table is None or index == table:
                            candidates.append((number, index, rows_found))
                if not candidates:
                    raise ValueError(
                        "No readable ruled table found. Export CSV/XLSX; scanned PDFs "
                        "and unruled tables are not supported"
                    )
                if len(candidates) != 1:
                    choices = ", ".join(f"page {p} table {t}" for p, t, _ in candidates)
                    raise ValueError(f"Choose page and table selectors: {choices}")
                rows = candidates[0][2]
        else:
            raise ValueError("Supported table files: .csv, .tsv, .xlsx, .pdf")
    except (ValueError, OSError):
        raise
    except ImportError as exc:
        raise ValueError(
            "Excel/PDF support requires: pip install efficientgraphbench[tables]"
        ) from exc
    except Exception as exc:
        raise ValueError(f"Cannot read {path.name}: {exc}") from exc
    rows = [["" if cell is None else str(cell) for cell in row] for row in rows]
    rows = rows[header_row - 1 :]
    if not rows:
        raise ValueError("No header at the selected header row")
    headers = rows[0]
    if (
        not headers
        or any(not name.strip() for name in headers)
        or len(set(headers)) != len(headers)
    ):
        raise ValueError("Table headers must be nonempty and unique")
    values = [row for row in rows[1:] if any(cell.strip() for cell in row)]
    if not values or any(len(row) != len(headers) for row in values):
        raise ValueError("Table must have data and the same number of cells in every row")
    return pd.DataFrame(values, columns=headers)
